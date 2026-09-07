import json
import sys
from pathlib import Path

from docx import Document


def paragraph_record(paragraph, location):
    return {
        "location": location,
        "style": paragraph.style.name if paragraph.style else None,
        "text": paragraph.text,
        "runs": [
            {
                "text": run.text,
                "bold": run.bold,
                "italic": run.italic,
                "underline": bool(run.underline) if run.underline is not None else None,
                "font": run.font.name,
                "size_pt": run.font.size.pt if run.font.size else None,
            }
            for run in paragraph.runs
        ],
    }


def main():
    src = Path(sys.argv[1])
    out = Path(sys.argv[2])
    doc = Document(src)
    payload = {
        "source": str(src),
        "sections": len(doc.sections),
        "paragraphs": [
            paragraph_record(p, f"body/p[{i}]") for i, p in enumerate(doc.paragraphs)
        ],
        "tables": [],
        "headers": [],
        "footers": [],
    }
    for ti, table in enumerate(doc.tables):
        rows = []
        for ri, row in enumerate(table.rows):
            cells = []
            for ci, cell in enumerate(row.cells):
                cells.append([
                    paragraph_record(p, f"table[{ti}]/row[{ri}]/cell[{ci}]/p[{pi}]")
                    for pi, p in enumerate(cell.paragraphs)
                ])
            rows.append(cells)
        payload["tables"].append(rows)
    for si, section in enumerate(doc.sections):
        payload["headers"].append([
            paragraph_record(p, f"section[{si}]/header/p[{i}]")
            for i, p in enumerate(section.header.paragraphs)
        ])
        payload["footers"].append([
            paragraph_record(p, f"section[{si}]/footer/p[{i}]")
            for i, p in enumerate(section.footer.paragraphs)
        ])
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
